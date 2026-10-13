"""Convierte el informe Markdown a Word (.docx).

Uso: python md_a_docx.py [entrada.md] [salida.docx]

Entrada por defecto:  Informe_CNN_Sismos.md
Salida por defecto:   Informe_CNN_Sismos.docx

Soporta el subconjunto de Markdown usado en el informe:
titulos (#, ##, ###), negritas, cursivas, codigo, subrayado <u>,
listas con viñetas y numeradas, tablas, imagenes y lineas en negrita
completas (portada).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

RAIZ = Path(__file__).resolve().parent
ENTRADA = RAIZ / "Informe_CNN_Sismos.md"
SALIDA = RAIZ / "Informe_CNN_Sismos.docx"

TOKEN = re.compile(
    r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|<u>.+?</u>)", re.DOTALL
)
TITULO = re.compile(r"^(#{1,4})\s+(.*)$")
NUMERADA = re.compile(r"^\d+\.\s+")


def configurar_estilos(doc: Document) -> None:
    """Times New Roman 12 con interlineado 1,5, tipico de trabajos academicos."""
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for nombre, tamano in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 13)):
        estilo = doc.styles[nombre]
        estilo.font.name = "Times New Roman"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.font.italic = False


def add_runs(parrafo, texto: str) -> None:
    """Escribe el texto de un parrafo aplicando formato inline."""
    texto = texto.replace("\\*", "*")
    for parte in TOKEN.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            run = parrafo.add_run(parte[2:-2])
            run.bold = True
        elif parte.startswith("<u>") and parte.endswith("</u>"):
            run = parrafo.add_run(parte[3:-4])
            run.underline = True
        elif parte.startswith("`") and parte.endswith("`"):
            parrafo.add_run(parte[1:-1])
        elif parte.startswith("*") and parte.endswith("*") and len(parte) > 2:
            run = parrafo.add_run(parte[1:-1])
            run.italic = True
        else:
            parrafo.add_run(parte)


def es_separador_tabla(linea: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linea.strip()))


def convertir(md_path: Path, docx_path: Path) -> dict:
    lineas = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configurar_estilos(doc)

    portada = True      # hasta el primer titulo, el texto se centra
    parrafo_texto: list[str] = []
    stats = {"parrafos": 0, "titulos": 0, "tablas": 0, "listas": 0, "imagenes": 0}

    def flush() -> None:
        nonlocal parrafo_texto
        if parrafo_texto:
            p = doc.add_paragraph()
            add_runs(p, " ".join(parrafo_texto))
            if portada:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            stats["parrafos"] += 1
            parrafo_texto = []

    i = 0
    while i < len(lineas):
        linea = lineas[i].rstrip()
        sin_espacios = linea.strip()

        if not sin_espacios:
            flush()
            i += 1
            continue

        # Titulo
        m = TITULO.match(sin_espacios)
        if m:
            flush()
            nivel = len(m.group(1))
            doc.add_heading(m.group(2), level=nivel)
            if portada:
                portada = False  # del titulo en adelante, cuerpo del informe
            stats["titulos"] += 1
            i += 1
            continue

        # Imagen
        if sin_espacios.startswith("!["):
            flush()
            ruta = re.sub(r"^!\[.*?\]\((.*?)\)$", r"\1", sin_espacios)
            if ruta.startswith("data:"):
                p = doc.add_paragraph("[logo embebido en el Markdown original]")
            else:
                img = RAIZ / ruta
                if img.exists():
                    doc.add_picture(str(img), width=Pt(240))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    stats["imagenes"] += 1
                else:
                    p = doc.add_paragraph(f"[Insertar imagen: {ruta}]")
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Tabla
        if sin_espacios.startswith("|"):
            flush()
            filas: list[list[str]] = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                bruta = lineas[i].strip()
                if not es_separador_tabla(bruta):
                    celdas = [c.strip() for c in bruta.strip("|").split("|")]
                    filas.append(celdas)
                i += 1
            if filas:
                n_cols = max(len(f) for f in filas)
                tabla = doc.add_table(rows=len(filas), cols=n_cols)
                tabla.style = "Table Grid"
                for r, fila in enumerate(filas):
                    for c in range(n_cols):
                        celda = tabla.cell(r, c)
                        celda.text = ""
                        add_runs(celda.paragraphs[0],
                                 fila[c] if c < len(fila) else "")
                        for run in celda.paragraphs[0].runs:
                            run.font.size = Pt(10)
                            run.font.name = "Times New Roman"
                        if r == 0:
                            for run in celda.paragraphs[0].runs:
                                run.bold = True
                stats["tablas"] += 1
                doc.add_paragraph()  # espacio tras la tabla
            continue

        # Lista con viñetas
        if sin_espacios.startswith("- "):
            flush()
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, sin_espacios[2:])
            stats["listas"] += 1
            i += 1
            continue

        # Lista numerada
        if NUMERADA.match(sin_espacios):
            flush()
            p = doc.add_paragraph(style="List Number")
            add_runs(p, NUMERADA.sub("", sin_espacios, count=1))
            stats["listas"] += 1
            i += 1
            continue

        # Linea entera en negrita (portada) o parrafo normal
        if re.fullmatch(r"\*\*.+\*\*", sin_espacios):
            flush()
            p = doc.add_paragraph()
            add_runs(p, sin_espacios)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            stats["parrafos"] += 1
            i += 1
            continue

        parrafo_texto.append(sin_espacios)
        i += 1

    flush()
    doc.save(docx_path)
    return stats


def main(argv: list[str]) -> int:
    md_path = Path(argv[1]) if len(argv) > 1 else ENTRADA
    docx_path = Path(argv[2]) if len(argv) > 2 else SALIDA
    if not md_path.is_file():
        print(f"No existe el archivo: {md_path}")
        return 1
    stats = convertir(md_path, docx_path)
    print(f"Convertido: {md_path.name} -> {docx_path}")
    print(f"  parrafos={stats['parrafos']} titulos={stats['titulos']} "
          f"tablas={stats['tablas']} listas={stats['listas']} "
          f"imagenes={stats['imagenes']}")
    # Verificacion: reabrir el .docx y comprobar que es legible
    doc = Document(docx_path)
    print(f"  verificacion: {len(doc.paragraphs)} parrafos, "
          f"{len(doc.tables)} tablas en el documento")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
